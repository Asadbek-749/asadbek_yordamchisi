"""Monthly Excel export of confirmed transactions (Uzbek headers, per-category totals)."""
from __future__ import annotations

import datetime
import io
from decimal import Decimal
from typing import Dict, List, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from discipline.db import get_session
from finance.models import FinancialTransaction
from finance.timeutil import month_label

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
BOLD = Font(bold=True)
MONEY_FORMAT = "#,##0.00"

TYPE_LABEL = {"INCOME": "Daromad", "EXPENSE": "Xarajat"}
STATUS_LABEL = {"CONFIRMED": "Tasdiqlangan", "EDITED": "Tahrirlangan"}


def _style_header(ws, row: int, count: int) -> None:
    for col in range(1, count + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _autosize(ws) -> None:
    for idx, column in enumerate(ws.columns, 1):
        longest = max((len(str(c.value)) for c in column if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(idx)].width = min(max(longest + 2, 10), 48)


async def load_month_transactions(owner_id: int, year: int, month: int) -> List[FinancialTransaction]:
    start = datetime.date(year, month, 1)
    end = datetime.date(year + 1, 1, 1) if month == 12 else datetime.date(year, month + 1, 1)
    async with get_session() as session:
        q = (
            select(FinancialTransaction)
            .where(
                FinancialTransaction.owner_id == owner_id,
                FinancialTransaction.transaction_date >= start,
                FinancialTransaction.transaction_date < end,
                FinancialTransaction.status.in_(["CONFIRMED", "EDITED"]),
            )
            .order_by(FinancialTransaction.transaction_date, FinancialTransaction.id)
        )
        return list((await session.execute(q)).scalars().all())


def build_month_workbook(txs: List[FinancialTransaction], year: int, month: int) -> bytes:
    """Render the workbook: sheet 1 = every transaction, 2 = per category, 3 = totals."""
    wb = Workbook()

    ws = wb.active
    ws.title = "Tranzaksiyalar"
    headers = ["Sana", "Turi", "Kategoriya", "Summa", "Valyuta", "Izoh", "Manba", "Holat"]
    ws.append(headers)
    _style_header(ws, 1, len(headers))
    for tx in txs:
        ws.append([
            tx.transaction_date,
            TYPE_LABEL.get(tx.type, tx.type),
            tx.category_name or "Boshqa",
            tx.amount,
            tx.currency,
            tx.description or "",
            tx.source,
            STATUS_LABEL.get(tx.status, tx.status),
        ])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        row[0].number_format = "DD.MM.YYYY"
        row[3].number_format = MONEY_FORMAT
    ws.freeze_panes = "A2"
    _autosize(ws)

    by_cat: Dict[Tuple[str, str, str], Decimal] = {}
    totals: Dict[Tuple[str, str], Decimal] = {}
    for tx in txs:
        by_cat[(tx.type, tx.category_name or "Boshqa", tx.currency)] = (
            by_cat.get((tx.type, tx.category_name or "Boshqa", tx.currency), Decimal("0")) + tx.amount
        )
        totals[(tx.currency, tx.type)] = totals.get((tx.currency, tx.type), Decimal("0")) + tx.amount

    ws2 = wb.create_sheet("Kategoriya bo'yicha")
    ws2.append(["Turi", "Kategoriya", "Valyuta", "Jami summa"])
    _style_header(ws2, 1, 4)
    for (tx_type, cat, cur), amount in sorted(by_cat.items(), key=lambda kv: (kv[0][0], kv[0][2], -kv[1])):
        ws2.append([TYPE_LABEL.get(tx_type, tx_type), cat, cur, amount])
    for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row):
        row[3].number_format = MONEY_FORMAT
    ws2.freeze_panes = "A2"
    _autosize(ws2)

    ws3 = wb.create_sheet("Xulosa")
    ws3.append([f"Moliyaviy xulosa — {month_label(year, month)}"])
    ws3["A1"].font = Font(bold=True, size=13)
    ws3.append([])
    ws3.append(["Valyuta", "Daromad", "Xarajat", "Sof balans"])
    _style_header(ws3, 3, 4)
    currencies = sorted({cur for cur, _ in totals}) or ["UZS"]
    for cur in currencies:
        inc = totals.get((cur, "INCOME"), Decimal("0"))
        exp = totals.get((cur, "EXPENSE"), Decimal("0"))
        ws3.append([cur, inc, exp, inc - exp])
    for row in ws3.iter_rows(min_row=4, max_row=ws3.max_row):
        for cell in row[1:]:
            cell.number_format = MONEY_FORMAT
    ws3.append([])
    ws3.append([f"Tranzaksiyalar soni: {len(txs)}"])
    _autosize(ws3)
    ws3.column_dimensions["A"].width = 16

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def export_month(owner_id: int, year: int, month: int) -> bytes:
    return build_month_workbook(await load_month_transactions(owner_id, year, month), year, month)
