"""Regression tests for the audit-and-hardening pass (extractor, brain, time zone, scheduler, db, currency)."""
from __future__ import annotations

import datetime
from decimal import Decimal
from unittest.mock import AsyncMock, MagicMock
from zoneinfo import ZoneInfo

import pytest
import pytest_asyncio

from discipline.db import close_db, init_db
from finance.extractor import FinanceExtractor
from finance.service import FinanceService
from nargis.brain import BrainIntent, NargisBrain

UTC = datetime.timezone.utc


def _txs(text: str):
    res = FinanceExtractor().rule_based_extract(text)
    return [(t.type, t.amount, t.category) for t in res.transactions]


# ── a) per-part income / expense detection ─────────────────────────────────

def test_mixed_income_and_expenses_are_classified_per_part():
    assert _txs("Taksi 15 ming, tushlik 45 ming, oylik 3 million oldim") == [
        ("EXPENSE", Decimal("15000"), "TRANSPORT"),
        ("EXPENSE", Decimal("45000"), "FOOD"),
        ("INCOME", Decimal("3000000"), "MAIN_JOB"),
    ]


def test_client_payment_then_internet_expense():
    got = _txs("Mijozdan 500 ming keldi, internetga 80 ming to'ladim")
    assert [(t, a) for t, a, _ in got] == [("INCOME", Decimal("500000")), ("EXPENSE", Decimal("80000"))]
    assert got[1][2] == "INTERNET"


def test_amountless_income_part_carries_meaning_to_next_part():
    assert _txs("Oylik oldim, 3 million") == [("INCOME", Decimal("3000000"), "MAIN_JOB")]


# ── b) dates, years, times, phones are not money ───────────────────────────

@pytest.mark.parametrize(
    "text",
    [
        "2026-yil 5-noyabr",
        "05.11.2026",
        "14:00",
        "+998 90 123 45 67",
        "998901234567",
        "Uchrashuv 2026-yil 5-noyabr soat 14:00 da",
        "Telefon +998 90 123 45 67",
    ],
)
def test_non_money_numbers_are_ignored(text):
    assert _txs(text) == []


def test_real_amounts_still_work():
    assert _txs("150 000 000 so'm")[0][1] == Decimal("150000000")
    assert _txs("Taksi 15000") == [("EXPENSE", Decimal("15000"), "TRANSPORT")]


# ── c) "Reja:" / "Maqsad:" win over money analysis ────────────────────────

def test_plan_prefix_is_not_an_expense():
    brain = NargisBrain(owner_id=1)
    assert brain.analyze_intent("Reja: 1) dars 2) 2026 hisobot", True).intent == BrainIntent.PLAN_CREATE
    assert brain.analyze_intent("Reja: bugun 500 ming sarflamayman", True).intent == BrainIntent.PLAN_CREATE


def test_goal_prefix_is_not_an_expense():
    brain = NargisBrain(owner_id=1)
    assert brain.analyze_intent("Maqsad: 5 million jamg'arish", True).intent == BrainIntent.GOAL_CREATE


# ── d) Tashkent time zone ──────────────────────────────────────────────────

def test_today_local_uses_tashkent(monkeypatch):
    from finance import timeutil

    # 20:00 UTC on 30 Oct is 01:00 on 31 Oct in Tashkent (UTC+5)
    monkeypatch.setattr(timeutil, "_utcnow", lambda: datetime.datetime(2026, 10, 30, 20, 0, tzinfo=UTC))
    assert timeutil.today_local() == datetime.date(2026, 10, 31)


def test_late_night_expense_lands_on_local_today(monkeypatch):
    from finance import timeutil

    monkeypatch.setattr(timeutil, "_utcnow", lambda: datetime.datetime(2026, 10, 30, 20, 0, tzinfo=UTC))
    res = FinanceExtractor().rule_based_extract("Taksi 15 ming")
    assert res.transactions[0].transaction_date == datetime.date(2026, 10, 31)
    res = FinanceExtractor().rule_based_extract("Kecha taksi 15 ming")
    assert res.transactions[0].transaction_date == datetime.date(2026, 10, 30)


# ── e) reminder catch-up window ────────────────────────────────────────────

def test_is_reminder_due_window():
    from discipline.scheduler import is_reminder_due

    tz = ZoneInfo("Asia/Tashkent")

    def at(h, m):
        return datetime.datetime(2026, 10, 30, h, m, tzinfo=tz)

    assert is_reminder_due(at(7, 0), "07:00")
    assert is_reminder_due(at(8, 29), "07:00")
    assert not is_reminder_due(at(8, 30), "07:00")
    assert not is_reminder_due(at(6, 59), "07:00")
    assert not is_reminder_due(at(7, 0), "garbage")


@pytest_asyncio.fixture
async def db_env(tmp_path):
    await init_db(data_dir=tmp_path)
    yield
    await close_db()


@pytest.mark.asyncio
async def test_morning_plan_catches_up_after_restart_and_is_sent_once(db_env):
    from discipline.scheduler import DisciplineScheduler

    bot = MagicMock()
    bot.send_message = AsyncMock()
    service = MagicMock()
    service.list_today_tasks = AsyncMock(return_value=[])
    sched = DisciplineScheduler(bot, service, owner_id=42)
    sched.proactive_engine = None

    # 07:45 local (02:45 UTC): the 07:00 minute was missed (bot was down)
    now_utc = datetime.datetime(2026, 10, 30, 2, 45, tzinfo=UTC)
    await sched._check_morning_and_evening_schedules(now_utc)
    await sched._check_morning_and_evening_schedules(now_utc + datetime.timedelta(minutes=1))

    morning = [c for c in bot.send_message.await_args_list if "Xayrli tong" in str(c)]
    assert len(morning) == 1


# ── f) PostgreSQL-safe schema extension ────────────────────────────────────

def test_schema_alterations_are_dialect_aware():
    from discipline.db import build_schema_alterations

    pg = build_schema_alterations("postgresql")
    assert pg and all("ADD COLUMN IF NOT EXISTS" in s for s in pg)
    assert any("BOOLEAN DEFAULT FALSE" in s for s in pg)
    assert not any("DEFAULT 0" in s and "BOOLEAN" in s for s in pg)

    lite = build_schema_alterations("sqlite")
    assert any("BOOLEAN DEFAULT 0" in s for s in lite)
    assert not any("IF NOT EXISTS" in s for s in lite)


# ── g) currency handling ───────────────────────────────────────────────────

def test_dollar_amounts_are_kept_as_usd():
    res = FinanceExtractor().rule_based_extract("Tushlik 100 dollar")
    tx = res.transactions[0]
    assert tx.amount == Decimal("100")
    assert tx.currency == "USD"
    res = FinanceExtractor().rule_based_extract("Kitob 25$ ")
    assert res.transactions[0].currency == "USD"
    assert FinanceExtractor().rule_based_extract("Tushlik 45 ming").transactions[0].currency == "UZS"


@pytest.mark.asyncio
async def test_reports_do_not_mix_currencies(db_env):
    from finance.extractor import ExtractedTransactionItem

    svc = FinanceService()
    today = datetime.date(2026, 10, 30)
    items = [
        ExtractedTransactionItem("EXPENSE", Decimal("45000"), "FOOD", "Tushlik", "UZS", today),
        ExtractedTransactionItem("EXPENSE", Decimal("100"), "FOOD", "Lunch", "USD", today),
    ]
    created = await svc.create_pending_transactions(7, items)
    for tx in created:
        await svc.confirm_transaction(tx.id, 7)

    daily = await svc.get_daily_summary(7, today)
    assert daily["total_expense"] == Decimal("45000")
    assert daily["by_currency"]["USD"]["expense"] == Decimal("100")

    monthly = await svc.get_monthly_summary(7, 2026, 10)
    assert monthly["total_expense"] == Decimal("45000")
    assert monthly["by_currency"]["USD"]["expense"] == Decimal("100")


@pytest.mark.asyncio
async def test_usd_rate_is_configurable(db_env):
    svc = FinanceService()
    assert await svc.get_usd_rate(7) is None
    await svc.set_usd_rate(7, Decimal("12650"))
    assert await svc.get_usd_rate(7) == Decimal("12650")
    await svc.set_usd_rate(7, Decimal("12700"))
    assert await svc.get_usd_rate(7) == Decimal("12700")


# ── voice monthly query crashed (get_monthly_summary was called with dates) ─

@pytest.mark.asyncio
async def test_voice_monthly_query_works(db_env):
    from nargis.voice import VoiceIntelligenceEngine

    engine = VoiceIntelligenceEngine(owner_id=1)
    res = await engine.process_voice_action("Bu oy qancha pul sarfladim?", user_id=1)
    assert res["success"] is True
    assert "Oylik" in res["response_text"]


# ── feature 1: AI fallback for money talk the rules do not understand ──────

class _JsonAI:
    def __init__(self, payload: str) -> None:
        self.payload = payload
        self.calls = 0

    async def generate(self, system_prompt: str, user_prompt: str) -> str:
        self.calls += 1
        return self.payload


@pytest.mark.asyncio
async def test_ai_fallback_used_only_when_rules_fail_and_text_is_financial():
    payload = '{"transactions": [{"type": "EXPENSE", "amount": 50000, "category": "DEBT", "currency": "UZS", "description": "Qarz"}], "ambiguous": false}'
    ai = _JsonAI(payload)
    ex = FinanceExtractor(ai)

    res = await ex.extract("Do'stimga ellik ming so'm qarz berdim")
    assert ai.calls == 1
    assert [(t.type, t.amount, t.category) for t in res.transactions] == [("EXPENSE", Decimal("50000"), "DEBT")]

    # rules succeed -> no AI call
    await ex.extract("Taksi 15 ming")
    assert ai.calls == 1
    # not financial -> no AI call
    res = await ex.extract("Salom, ishlar qalay bugun")
    assert ai.calls == 1 and not res.transactions


@pytest.mark.asyncio
async def test_ai_fallback_output_is_sanitised():
    payload = '{"transactions": [{"type": "TRANSFER", "amount": 10, "category": "X"}, {"type": "income", "amount": "2500000", "category": "NOPE", "currency": "usd"}]}'
    res = await FinanceExtractor(_JsonAI(payload)).extract("Menga ikki yarim million dollar keldi")
    assert [(t.type, t.amount, t.category, t.currency) for t in res.transactions] == [("INCOME", Decimal("2500000"), "OTHER", "USD")]


# ── feature 2: editing a pending record ────────────────────────────────────

@pytest.mark.asyncio
async def test_edit_pending_amount_and_category(db_env):
    from finance.extractor import ExtractedTransactionItem
    from finance.keyboards import category_choice_keyboard, pending_transaction_keyboard

    svc = FinanceService()
    item = ExtractedTransactionItem("EXPENSE", Decimal("45000"), "FOOD", "Tushlik", "UZS", datetime.date(2026, 10, 30))
    tx = (await svc.create_pending_transactions(7, [item]))[0]

    assert any(b.callback_data.startswith("fin:edit:") for b in pending_transaction_keyboard(tx.id).inline_keyboard[0])

    edited = await svc.edit_transaction(tx.id, 7, new_amount=Decimal("60000"), new_category_key="TRANSPORT")
    assert edited.amount == Decimal("60000")
    assert edited.status == "PENDING"  # still waits for confirmation
    assert "Transport" in edited.category_name

    cats = await svc.list_categories("EXPENSE")
    kb = category_choice_keyboard(tx.id, cats)
    assert all(len(b.callback_data.encode()) <= 64 for row in kb.inline_keyboard for b in row)

    # only the owner can load the record
    assert await svc.get_transaction(tx.id, 8) is None
    assert (await svc.get_transaction(tx.id, 7)).id == tx.id


@pytest.mark.asyncio
async def test_cancelled_record_cannot_be_revived_by_old_button(db_env):
    from finance.extractor import ExtractedTransactionItem

    svc = FinanceService()
    item = ExtractedTransactionItem("EXPENSE", Decimal("1000"), "FOOD", "x", "UZS", datetime.date(2026, 10, 30))
    tx = (await svc.create_pending_transactions(7, [item]))[0]
    await svc.cancel_transaction(tx.id, 7)
    assert await svc.confirm_transaction(tx.id, 7) is None
    assert await svc.edit_transaction(tx.id, 7, new_amount=Decimal("5")) is None


# ── feature 3: /export ─────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_export_month_workbook(db_env):
    import io

    from openpyxl import load_workbook

    from finance.export import export_month
    from finance.extractor import ExtractedTransactionItem

    svc = FinanceService()
    d = datetime.date(2026, 10, 12)
    items = [
        ExtractedTransactionItem("EXPENSE", Decimal("15000"), "TRANSPORT", "Taksi", "UZS", d),
        ExtractedTransactionItem("EXPENSE", Decimal("45000"), "FOOD", "Tushlik", "UZS", d),
        ExtractedTransactionItem("EXPENSE", Decimal("25000"), "FOOD", "Kechki ovqat", "UZS", d),
        ExtractedTransactionItem("INCOME", Decimal("3000000"), "MAIN_JOB", "Oylik", "UZS", d),
        ExtractedTransactionItem("EXPENSE", Decimal("100"), "SHOPPING", "Kitob", "USD", d),
        ExtractedTransactionItem("EXPENSE", Decimal("999"), "FOOD", "Boshqa oy", "UZS", datetime.date(2026, 9, 30)),
    ]
    created = await svc.create_pending_transactions(7, items)
    pending_only = created[0]
    for tx in created[1:]:
        await svc.confirm_transaction(tx.id, 7)  # the first one stays PENDING

    wb = load_workbook(io.BytesIO(await export_month(7, 2026, 10)))
    assert wb.sheetnames == ["Tranzaksiyalar", "Kategoriya bo'yicha", "Xulosa"]

    rows = list(wb["Tranzaksiyalar"].iter_rows(values_only=True))
    assert rows[0] == ("Sana", "Turi", "Kategoriya", "Summa", "Valyuta", "Izoh", "Manba", "Holat")
    assert len(rows) == 1 + 4  # confirmed this-month records only (no pending, no September)

    cat_rows = {(r[0], r[1], r[2]): r[3] for r in wb["Kategoriya bo'yicha"].iter_rows(min_row=2, values_only=True)}
    food = [v for (t, c, cur), v in cat_rows.items() if "Ovqat" in c and cur == "UZS"]
    assert food and float(food[0]) == 70000.0

    summary = {r[0]: r[1:] for r in wb["Xulosa"].iter_rows(min_row=4, max_row=5, values_only=True)}
    assert float(summary["UZS"][0]) == 3000000.0 and float(summary["UZS"][1]) == 70000.0
    assert float(summary["USD"][1]) == 100.0
    assert pending_only.id  # silence unused warning
